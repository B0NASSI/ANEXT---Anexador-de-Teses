# -*- coding: utf-8 -*-
"""
Montagem de cada variante de página a partir de modelo.docx: substitui o
placeholder {{TITULO}} (capa geral = título da tese; página de segurado =
nome do segurado, centralizado pelo próprio estilo do parágrafo do modelo)
e substitui o placeholder {{TABELA}} por uma tabela real, clonando a
aparência (mesclagens, cores de cabeçalho, bordas, fontes) da tabela de
origem quando ela existe (importação de .docx). Sem tabela de origem
(texto colado), aplica um estilo padrão na paleta da Rodriguez & Sousa.

Os parágrafos são varridos em todo o corpo do documento (`iter`), não só
`doc.paragraphs`, porque um {{TITULO}} de capa costuma estar dentro de uma
caixa de texto — que python-docx não expõe na API de alto nível, mas que
ainda é um `<w:p>` alcançável percorrendo a árvore XML.
"""

import copy
import re
import unicodedata

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

MARCADOR_TITULO = "{{TITULO}}"
MARCADOR_TITULO_CAPA = "{{TITULO_CAPA}}"
MARCADOR_TOPICO = "{{TOPICO}}"
MARCADOR_TABELA = "{{TABELA}}"

COR_TABELA_DESTAQUE = "D5DCE4"


# ── Localização e substituição de texto em qualquer parágrafo do corpo ───

def _todos_paragrafos(doc: Document):
    for p_elem in doc.element.body.iter(qn("w:p")):
        yield Paragraph(p_elem, doc)
    for secao in doc.sections:
        for parte in (
            secao.header, secao.footer,
            secao.first_page_header, secao.first_page_footer,
            secao.even_page_header, secao.even_page_footer,
        ):
            for p in parte.paragraphs:
                yield p


def _substituir_no_paragrafo(paragrafo: Paragraph, marcador: str, novo_texto: str) -> bool:
    # Caso comum: o marcador inteiro está dentro de um único run — troca só
    # esse run, sem tocar nos vizinhos. Importante quando o parágrafo tem
    # elementos não textuais ao lado (ex.: uma quebra de página manual),
    # que o caminho abaixo (esvaziar runs[1:]) apagaria.
    for run in paragrafo.runs:
        if marcador in run.text:
            run.text = run.text.replace(marcador, novo_texto)
            return True

    # Caso raro: o marcador foi dividido entre vários runs (autocorreção do
    # Word, por exemplo) — junta o texto de todos, substitui, e concentra o
    # resultado no primeiro run, esvaziando os demais.
    texto_completo = "".join(r.text for r in paragrafo.runs)
    if marcador not in texto_completo:
        return False
    novo = texto_completo.replace(marcador, novo_texto)
    if paragrafo.runs:
        paragrafo.runs[0].text = novo
        for run in paragrafo.runs[1:]:
            run.text = ""
    else:
        paragrafo.add_run(novo)
    return True


def substituir_titulo(doc: Document, texto_novo: str, marcador: str = MARCADOR_TITULO) -> bool:
    encontrado = False
    for paragrafo in _todos_paragrafos(doc):
        if _substituir_no_paragrafo(paragrafo, marcador, texto_novo):
            encontrado = True
    return encontrado


def substituir_topico(doc: Document, numero: str) -> bool:
    return substituir_titulo(doc, numero, marcador=MARCADOR_TOPICO)


def localizar_paragrafo_marcador(doc: Document, marcador: str = MARCADOR_TABELA) -> Paragraph | None:
    for paragrafo in _todos_paragrafos(doc):
        if marcador in "".join(r.text for r in paragrafo.runs):
            return paragrafo
    return None


def remover_paragrafo(paragrafo: Paragraph) -> None:
    elemento = paragrafo._p
    elemento.getparent().remove(elemento)


def _remover_ate(doc: Document, elemento_final) -> None:
    """Remove todo o conteúdo do corpo do documento desde o início até
    `elemento_final`, inclusive — usado para apagar o bloco da capa (que vem
    antes) ao montar a página de tabela completa/segurado."""
    body = doc.element.body
    for filho in list(body):
        if filho.tag == qn("w:sectPr"):
            break
        body.remove(filho)
        if filho is elemento_final:
            break


def _remover_de(doc: Document, elemento_inicial) -> None:
    """Remove `elemento_inicial` e tudo que vem depois dele no corpo do
    documento, preservando o `sectPr` final (tamanho/margens da página) —
    usado para apagar o bloco de título+tabela ao montar só a capa."""
    body = doc.element.body
    removendo = False
    for filho in list(body):
        if filho is elemento_inicial:
            removendo = True
        if removendo:
            if filho.tag == qn("w:sectPr"):
                break
            body.remove(filho)


# ── Clonagem de aparência de célula/fonte a partir de uma tabela de origem ─

ATRIBUTOS_TEMA = (
    qn("w:themeColor"), qn("w:themeTint"), qn("w:themeShade"),
    qn("w:themeFill"), qn("w:themeFillTint"), qn("w:themeFillShade"),
)


def _remover_referencias_de_tema(elemento) -> None:
    """Cores por tema (`w:themeColor`/`w:themeFill`) resolvem contra o
    theme1.xml de CADA documento — um azul clarinho na tabela de origem
    pode virar outra cor (ou sumir) no modelo, que tem um tema diferente.
    Remove essas referências do elemento e de todos os descendentes,
    mantendo só a cor absoluta (`w:val`/`w:fill`, RGB explícito)."""
    for sub in elemento.iter():
        for atributo in ATRIBUTOS_TEMA:
            if atributo in sub.attrib:
                del sub.attrib[atributo]


def _clonar_elemento(elemento):
    clone = copy.deepcopy(elemento)
    _remover_referencias_de_tema(clone)
    return clone


def _clonar_aparencia_celula(origem_cell, destino_cell) -> None:
    tcPr_origem = origem_cell._tc.tcPr
    if tcPr_origem is None:
        return
    tcPr_destino = destino_cell._tc.get_or_add_tcPr()
    for tag in ("w:tcW", "w:shd", "w:tcBorders", "w:vAlign"):
        elemento = tcPr_origem.find(qn(tag))
        if elemento is None:
            continue
        existente = tcPr_destino.find(qn(tag))
        if existente is not None:
            tcPr_destino.remove(existente)
        tcPr_destino.append(_clonar_elemento(elemento))


def _clonar_largura_tabela(tabela_referencia, tabela_nova) -> None:
    """Clona a largura da tabela de origem (largura total + grade de colunas
    + tipo de ajuste fixo/automático) para a tabela nova. É proibido deixar
    o Word recalcular essas larguras — sem isso, `add_table()` cria colunas
    de largura igual, ignorando a proporção real da tabela original."""
    tblPr_origem = tabela_referencia._tbl.tblPr
    tblPr_novo = tabela_nova._tbl.tblPr
    if tblPr_origem is not None and tblPr_novo is not None:
        # w:tblLook precisa vir junto: são as flags de "primeira coluna"/
        # "primeira linha em destaque" (bandas). `add_table()` liga
        # firstColumn por padrão — sem clonar isso da origem, o Word aplica
        # uma formatação condicional de cor completamente diferente.
        for tag in ("w:tblW", "w:tblLayout", "w:tblLook"):
            elemento = tblPr_origem.find(qn(tag))
            if elemento is None:
                continue
            existente = tblPr_novo.find(qn(tag))
            if existente is not None:
                tblPr_novo.remove(existente)
            tblPr_novo.append(_clonar_elemento(elemento))

    grade_origem = tabela_referencia._tbl.find(qn("w:tblGrid"))
    grade_nova = tabela_nova._tbl.find(qn("w:tblGrid"))
    if grade_origem is not None and grade_nova is not None:
        grade_nova.getparent().replace(grade_nova, _clonar_elemento(grade_origem))


def _clonar_fonte_paragrafo(origem_paragrafo, destino_paragrafo) -> None:
    rpr_origem = origem_paragrafo.runs[0]._r.rPr if origem_paragrafo.runs else None
    destino_run = destino_paragrafo.runs[0] if destino_paragrafo.runs else destino_paragrafo.add_run("")
    if rpr_origem is not None:
        rpr_destino = destino_run._r.get_or_add_rPr()
        pai = rpr_destino.getparent()
        indice = list(pai).index(rpr_destino)
        pai.remove(rpr_destino)
        pai.insert(indice, _clonar_elemento(rpr_origem))
    destino_paragrafo.alignment = origem_paragrafo.alignment

    # espaçamento antes/depois do parágrafo (w:spacing) é o que dá a altura
    # visual da linha na tabela de origem — sem clonar isso, add_table()
    # cria linhas "encolhidas" (só a altura mínima do texto).
    pPr_origem = origem_paragrafo._p.pPr
    spacing_origem = pPr_origem.find(qn("w:spacing")) if pPr_origem is not None else None
    if spacing_origem is not None:
        pPr_destino = destino_paragrafo._p.get_or_add_pPr()
        existente = pPr_destino.find(qn("w:spacing"))
        if existente is not None:
            pPr_destino.remove(existente)
        pPr_destino.append(_clonar_elemento(spacing_origem))


def _sombrear_celula(celula, cor_hex: str) -> None:
    tcPr = celula._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = tcPr.makeelement(qn("w:shd"), {})
        tcPr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor_hex)


def _forcar_negrito(celula) -> None:
    for paragrafo in celula.paragraphs:
        for run in paragrafo.runs:
            run.font.bold = True


def _detectar_coluna_item(cabecalhos: list[str]) -> int | None:
    for i, cabecalho in enumerate(cabecalhos):
        if cabecalho.strip().lower() == "item":
            return i
    return None


def _clonar_altura_linha(origem_row, destino_row) -> None:
    """`add_table()` cria linhas com altura mínima (só o suficiente para o
    texto) — para a tabela sair do mesmo tamanho da tabela Word enviada,
    precisa herdar a altura explícita da linha de origem."""
    if origem_row.height is not None:
        destino_row.height = origem_row.height
    if origem_row.height_rule is not None:
        destino_row.height_rule = origem_row.height_rule


# ── Catálogo de larguras oficiais (TABELAS.docx) ──────────────────────────
#
# TABELAS.docx (na pasta modelo/) traz uma tabela-exemplo de cada tipo de
# tese, com as larguras de coluna oficiais do escritório. Na geração, o
# tipo da tabela importada é identificado pelos cabeçalhos e as larguras
# oficiais são aplicadas — as larguras do arquivo importado variam de
# usuário para usuário e por isso são ignoradas quando há tipo catalogado.

_CATALOGO_TABELAS: dict = {}


# variações de gênero/número do cabeçalho de nomes — EMPREGADA(S),
# EMPREGADO(A), EMPREGADOS (AS)... — todas equivalem a "empregados" na
# identificação do tipo de tabela no catálogo
_RE_EMPREGADOS = re.compile(r"empregad[oa]s?(\s*\(\s*[oa]s?\s*\))?")


def _normalizar_cabecalho(texto: str) -> str:
    sem_acento = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode("ascii")
    base = " ".join(sem_acento.lower().split())
    return _RE_EMPREGADOS.sub("empregados", base)


def _chave_cabecalhos(cabecalhos_visuais: list[str]) -> tuple:
    return tuple(_normalizar_cabecalho(c) for c in cabecalhos_visuais)


def _cabecalhos_visuais_de_linha(linha_word) -> list[str]:
    visuais = []
    vistos = set()
    for celula in linha_word.cells:
        if id(celula._tc) not in vistos:
            vistos.add(id(celula._tc))
            visuais.append(" ".join(p.text for p in celula.paragraphs).strip())
    return visuais


def carregar_catalogo_tabelas(caminho) -> None:
    """Lê o TABELAS.docx e indexa, por (cabeçalhos, nº de colunas da grade),
    as larguras oficiais de cada tipo de tabela. Chamado antes da geração;
    se o arquivo não existir, o catálogo fica vazio e nada muda."""
    global _CATALOGO_TABELAS
    _CATALOGO_TABELAS = {}
    try:
        doc = Document(str(caminho))
    except Exception:
        return
    for tabela_word in doc.tables:
        if not tabela_word.rows:
            continue
        grade = tabela_word._tbl.find(qn("w:tblGrid"))
        if grade is None:
            continue
        larguras = [int(c.get(qn("w:w")) or 0) for c in grade.findall(qn("w:gridCol"))]
        if not larguras or not all(larguras):
            continue
        visuais = _cabecalhos_visuais_de_linha(tabela_word.rows[0])
        chave = (_chave_cabecalhos(visuais), len(larguras))
        _CATALOGO_TABELAS.setdefault(chave, larguras)


def _larguras_do_catalogo(cabecalhos: list[str]) -> list[int] | None:
    if not _CATALOGO_TABELAS:
        return None
    spans = getattr(cabecalhos, "spans", None)
    if spans:
        visuais = [c for c, s in zip(cabecalhos, spans) if s != 0]
    else:
        visuais = list(cabecalhos)
    return _CATALOGO_TABELAS.get((_chave_cabecalhos(visuais), len(cabecalhos)))


def _aplicar_larguras_twips(tabela, larguras: list[int]) -> None:
    """Aplica larguras absolutas (twips) às colunas: grade, largura total,
    largura de cada célula (somando mesclagens horizontais) e layout fixo."""
    tbl = tabela._tbl
    grade = tbl.find(qn("w:tblGrid"))
    if grade is None:
        return
    colunas_grade = grade.findall(qn("w:gridCol"))
    if len(colunas_grade) != len(larguras):
        return
    for coluna, largura in zip(colunas_grade, larguras):
        coluna.set(qn("w:w"), str(largura))

    tblPr = tbl.tblPr
    for tag, atributos in (
        ("w:tblW", {"w:w": str(sum(larguras)), "w:type": "dxa"}),
        ("w:tblLayout", {"w:type": "fixed"}),
    ):
        elemento = tblPr.find(qn(tag))
        if elemento is None:
            elemento = tblPr.makeelement(qn(tag), {})
            tblPr.append(elemento)
        for chave, valor in atributos.items():
            elemento.set(qn(chave), valor)

    for linha in tabela.rows:
        indice = 0
        for tc in linha._tr.findall(qn("w:tc")):
            tcPr = tc.find(qn("w:tcPr"))
            if tcPr is None:
                tcPr = tc.makeelement(qn("w:tcPr"), {})
                tc.insert(0, tcPr)
            grid_span = tcPr.find(qn("w:gridSpan"))
            span = int(grid_span.get(qn("w:val"))) if grid_span is not None else 1
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = tcPr.makeelement(qn("w:tcW"), {})
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(sum(larguras[indice:indice + span])))
            tcW.set(qn("w:type"), "dxa")
            indice += span


def _centralizar_tabela(tabela) -> None:
    tblPr = tabela._tbl.tblPr
    jc = tblPr.find(qn("w:jc"))
    if jc is None:
        jc = tblPr.makeelement(qn("w:jc"), {})
        tblPr.append(jc)
    jc.set(qn("w:val"), "center")


# ── Cópia fiel de linhas (w:tr) e formas flutuantes do documento original ──
#
# A tabela importada é reproduzida copiando (deepcopy) os nós w:tr inteiros
# do documento de origem — preserva formatação inline (negrito/cor aplicados
# em runs individuais), mesclagens e as formas vetoriais (setas) ancoradas
# dentro das células. Formas ancoradas fora da região copiada (no parágrafo
# após a tabela, ou na primeira linha do segurado seguinte, flutuando por
# cima da linha anterior) são trazidas num parágrafo logo após a tabela
# nova, mantendo as coordenadas de ancoragem; o Word renderiza o documento
# completo sem reposicionamento manual.

_NS_MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
_NS_WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
_TAG_ALTERNATE_CONTENT = _NS_MC + "AlternateContent"


def _ancora_da_forma(alternate_content):
    for anchor in alternate_content.iter(_NS_WP + "anchor"):
        return anchor
    return None


def _pos_offset(anchor, tag: str):
    """(elemento wp:positionH/V, offset em EMU ou None)."""
    pos = anchor.find(_NS_WP + tag)
    if pos is None:
        return None, None
    off = pos.find(_NS_WP + "posOffset")
    try:
        return pos, int(off.text)
    except (AttributeError, TypeError, ValueError):
        return pos, None


def _formas_flutuando_acima(elemento):
    """Formas (mc:AlternateContent) ancoradas em `elemento` cujo offset
    vertical é negativo — flutuam por cima do conteúdo ANTERIOR à âncora."""
    formas = []
    for ac in elemento.iter(_TAG_ALTERNATE_CONTENT):
        anchor = _ancora_da_forma(ac)
        if anchor is None:
            continue
        _, offset_v = _pos_offset(anchor, "positionV")
        if offset_v is not None and offset_v < 0:
            formas.append(ac)
    return formas


def _tcs_com_coluna_grade(tr):
    """[(índice da coluna da grade onde a célula começa, w:tc), ...]"""
    resultado = []
    coluna = 0
    for tc in tr.findall(qn("w:tc")):
        tcPr = tc.find(qn("w:tcPr"))
        span = 1
        if tcPr is not None:
            grid_span = tcPr.find(qn("w:gridSpan"))
            if grid_span is not None:
                span = int(grid_span.get(qn("w:val")) or 1)
        resultado.append((coluna, tc))
        coluna += span
    return resultado


def _eh_vmerge_continue(tc) -> bool:
    tcPr = tc.find(qn("w:tcPr"))
    if tcPr is None:
        return False
    vmerge = tcPr.find(qn("w:vMerge"))
    return vmerge is not None and (vmerge.get(qn("w:val")) or "continue") == "continue"


def _definir_vmerge_restart(tc) -> None:
    tcPr = tc.find(qn("w:tcPr"))
    if tcPr is None:
        tcPr = tc.makeelement(qn("w:tcPr"), {})
        tc.insert(0, tcPr)
    vmerge = tcPr.find(qn("w:vMerge"))
    if vmerge is None:
        vmerge = tcPr.makeelement(qn("w:vMerge"), {})
        tcPr.append(vmerge)
    vmerge.set(qn("w:val"), "restart")


def _promover_vmerges_orfaos(tr_novo, trs_ref, idx_origem: int) -> None:
    """Quando a linha copiada era continuação de mesclagem vertical mas a
    linha de cima NÃO veio junto (página por segurado), a célula-continuação
    ficaria pendurada na linha errada. Substitui pela cópia da célula que
    iniciou a mesclagem (com o texto e a formatação originais), marcada
    como restart."""
    for coluna, tc in _tcs_com_coluna_grade(tr_novo):
        if not _eh_vmerge_continue(tc):
            continue
        origem = None
        for i in range(idx_origem - 1, 0, -1):
            for col_ref, tc_ref in _tcs_com_coluna_grade(trs_ref[i]):
                if col_ref == coluna:
                    if not _eh_vmerge_continue(tc_ref):
                        origem = tc_ref
                    break
            if origem is not None:
                break
        if origem is not None:
            tc_novo = _clonar_elemento(origem)
            _definir_vmerge_restart(tc_novo)
            tc.getparent().replace(tc, tc_novo)
        else:
            _definir_vmerge_restart(tc)


def _remover_formas_acima(tr_novo) -> None:
    """Remove formas ancoradas nesta linha que flutuam por cima da linha
    anterior do documento original — quando essa linha anterior não foi
    copiada, a forma apareceria sobre o conteúdo errado (ex.: cabeçalho)."""
    for ac in _formas_flutuando_acima(tr_novo):
        ac.getparent().remove(ac)


def _chave_numeracao(paragrafo_elem):
    """(numId, nível) do parágrafo, se ele usa numeração automática do Word."""
    pPr = paragrafo_elem.find(qn("w:pPr"))
    if pPr is None:
        return None
    numPr = pPr.find(qn("w:numPr"))
    if numPr is None:
        return None
    num_id = numPr.find(qn("w:numId"))
    valor = num_id.get(qn("w:val")) if num_id is not None else None
    if valor in (None, "0"):
        return None
    ilvl = numPr.find(qn("w:ilvl"))
    return (valor, ilvl.get(qn("w:val")) if ilvl is not None else "0")


def _resolver_numeracao_automatica(tbl_ref) -> dict:
    """Resolve, na ordem do documento original, o número de cada parágrafo
    com numeração automática dentro da tabela — permite materializar o
    número certo mesmo quando a página copia só as linhas de um segurado
    (o item 4 continua sendo "4")."""
    contadores: dict = {}
    numeros: dict = {}
    for p in tbl_ref.iter(qn("w:p")):
        chave = _chave_numeracao(p)
        if chave is None:
            continue
        contadores[chave] = contadores.get(chave, 0) + 1
        numeros[p] = str(contadores[chave])
    return numeros


def _materializar_numeracao(tr_novo, tr_origem, numeros: dict) -> None:
    """Troca a numeração automática das células copiadas por texto literal
    (só o número, sem o "." que o formato de lista do Word acrescenta)."""
    pares = zip(
        [p for p in tr_origem.iter(qn("w:p")) if _chave_numeracao(p) is not None],
        [p for p in tr_novo.iter(qn("w:p")) if _chave_numeracao(p) is not None],
    )
    for p_origem, p_novo in pares:
        numero = numeros.get(p_origem)
        if numero is None:
            continue
        pPr = p_novo.find(qn("w:pPr"))
        pPr.remove(pPr.find(qn("w:numPr")))
        run = OxmlElement("w:r")
        rPr_marca = pPr.find(qn("w:rPr"))
        if rPr_marca is not None:
            run.append(copy.deepcopy(rPr_marca))
        t = OxmlElement("w:t")
        t.text = numero
        run.append(t)
        p_novo.insert(list(p_novo).index(pPr) + 1, run)


def _construir_tabela_por_copia(doc: Document, tabela_referencia, linhas_dados) -> Table:
    tbl_ref = tabela_referencia._tbl
    trs_ref = tbl_ref.findall(qn("w:tr"))
    indice_de = {tr: i for i, tr in enumerate(trs_ref)}
    numeracao = _resolver_numeracao_automatica(tbl_ref)

    novo = copy.deepcopy(tbl_ref)
    for tr in novo.findall(qn("w:tr")):
        novo.remove(tr)
    cabecalho_novo = copy.deepcopy(trs_ref[0])
    _materializar_numeracao(cabecalho_novo, trs_ref[0], numeracao)
    novo.append(cabecalho_novo)

    idx_anterior = 0
    for linha in linhas_dados:
        idx = indice_de[linha.tr]
        tr_novo = copy.deepcopy(linha.tr)
        _materializar_numeracao(tr_novo, linha.tr, numeracao)
        if idx != idx_anterior + 1:
            _promover_vmerges_orfaos(tr_novo, trs_ref, idx)
            _remover_formas_acima(tr_novo)
        idx_anterior = idx
        novo.append(tr_novo)

    _remover_referencias_de_tema(novo)
    doc.element.body.append(novo)
    tabela = Table(novo, doc)
    try:
        if tabela_referencia.style is not None:
            tabela.style = tabela_referencia.style.name
    except Exception:
        pass
    return tabela


def _atualizar_posicao_h(alternate_content, novo_offset_emu: int) -> None:
    """Define a posição horizontal absoluta (relativa à coluna de texto do
    corpo) na âncora moderna (wp:anchor) e no fallback VML (margin-left)."""
    anchor = _ancora_da_forma(alternate_content)
    if anchor is None:
        return
    pos, _ = _pos_offset(anchor, "positionH")
    if pos is not None:
        pos.set("relativeFrom", "column")
        off = pos.find(_NS_WP + "posOffset")
        if off is None:
            for filho in list(pos):
                pos.remove(filho)
            off = pos.makeelement(_NS_WP + "posOffset", {})
            pos.append(off)
        off.text = str(novo_offset_emu)
    margem_pt = novo_offset_emu / 12700.0
    for shape in alternate_content.iter():
        estilo = shape.get("style")
        if estilo and "margin-left" in estilo:
            shape.set("style", re.sub(r"margin-left:[^;]*", f"margin-left:{margem_pt:.2f}pt", estilo))


def _paragrafo_com_formas(formas) -> "OxmlElement":
    p = OxmlElement("w:p")
    run = OxmlElement("w:r")
    p.append(run)
    for forma in formas:
        run.append(forma)
    return p


def _anexar_setas_flutuantes(doc: Document, tabela: Table, tabela_referencia, linhas_dados) -> None:
    """Coleta as formas flutuantes que pertencem à última linha copiada mas
    estavam ancoradas FORA dela no original — no parágrafo logo após a
    tabela (último segurado) ou na primeira linha do segurado seguinte — e
    as insere num parágrafo após a tabela nova. O offset vertical original é
    mantido (negativo, flutuando sobre a última linha); o horizontal, quando
    a âncora era relativa à coluna de uma célula, é recalculado para a
    posição absoluta da mesma coluna na tabela nova (centralizada)."""
    tbl_ref = tabela_referencia._tbl
    trs_ref = tbl_ref.findall(qn("w:tr"))
    try:
        ultimo_idx = trs_ref.index(linhas_dados[-1].tr)
    except (ValueError, AttributeError):
        return

    paragrafos = []
    if ultimo_idx == len(trs_ref) - 1:
        # última linha do original incluída: as formas do parágrafo seguinte
        # à tabela original valem como estão (mesmas coordenadas)
        for irmao in tbl_ref.itersiblings():
            if irmao.tag in (qn("w:tbl"), qn("w:sectPr")):
                break
            if irmao.tag == qn("w:p") and _formas_flutuando_acima(irmao):
                clone = _clonar_elemento(irmao)
                paragrafos.append(clone)
    else:
        grade = tabela._tbl.find(qn("w:tblGrid"))
        larguras = [int(c.get(qn("w:w")) or 0) for c in grade.findall(qn("w:gridCol"))] if grade is not None else []
        seccao = doc.sections[0]
        util_twips = int((seccao.page_width - seccao.left_margin - seccao.right_margin) / 635)
        # pode ser NEGATIVO: tabela mais larga que a coluna de texto,
        # centralizada, começa à esquerda da margem — o offset da âncora
        # (relativo à coluna de texto) precisa acompanhar
        esquerda_twips = (util_twips - sum(larguras)) // 2

        formas_novas = []
        for coluna, tc in _tcs_com_coluna_grade(trs_ref[ultimo_idx + 1]):
            for ac in _formas_flutuando_acima(tc):
                anchor = _ancora_da_forma(ac)
                _, offset_h = _pos_offset(anchor, "positionH")
                clone = _clonar_elemento(ac)
                if larguras and offset_h is not None:
                    # 108 twips = margem interna lateral padrão da célula
                    x_celula_twips = esquerda_twips + sum(larguras[:coluna]) + 108
                    _atualizar_posicao_h(clone, x_celula_twips * 635 + offset_h)
                formas_novas.append(clone)
        if formas_novas:
            paragrafos.append(_paragrafo_com_formas(formas_novas))

    if paragrafos:
        referencia = tabela._tbl
        for p in paragrafos:
            referencia.addnext(p)
            referencia = p
        tabela._paragrafos_setas = paragrafos


# ── Construção e inserção da tabela ───────────────────────────────────────

def _aplicar_estilo_padrao_grade(tabela) -> None:
    """Estilo aplicado quando não há tabela de origem para clonar (texto
    colado sem formatação): grade completa com bordas finas pretas e todo
    o conteúdo centralizado, seguindo o padrão das tabelas FAP."""
    tblPr = tabela._tbl.tblPr
    existente = tblPr.find(qn("w:tblBorders"))
    if existente is not None:
        tblPr.remove(existente)
    bordas = tblPr.makeelement(qn("w:tblBorders"), {})
    for lado in ("w:top", "w:left", "w:bottom", "w:right", "w:insideH", "w:insideV"):
        borda = bordas.makeelement(qn(lado), {})
        borda.set(qn("w:val"), "single")
        borda.set(qn("w:sz"), "4")
        borda.set(qn("w:space"), "0")
        borda.set(qn("w:color"), "000000")
        bordas.append(borda)
    tblPr.append(bordas)

    for linha in tabela.rows:
        for celula in linha.cells:
            celula.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragrafo in celula.paragraphs:
                paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _aplicar_larguras_colunas(doc: Document, tabela, proporcoes: list[float]) -> None:
    """Aplica larguras de coluna personalizadas (proporções que somam 1)
    à tabela reconstruída: grade, largura total, layout fixo e a largura de
    cada célula (somando as colunas cobertas por mesclagens horizontais)."""
    seccao = doc.sections[0]
    # a subtração de Length devolve int em EMU; 635 EMU = 1 twip
    total_twips = int((seccao.page_width - seccao.left_margin - seccao.right_margin) / 635)
    larguras = [max(100, int(total_twips * p)) for p in proporcoes]

    tbl = tabela._tbl
    grade = tbl.find(qn("w:tblGrid"))
    if grade is None:
        return
    colunas_grade = grade.findall(qn("w:gridCol"))
    if len(colunas_grade) != len(larguras):
        return
    for coluna, largura in zip(colunas_grade, larguras):
        coluna.set(qn("w:w"), str(largura))

    tblPr = tbl.tblPr
    for tag, atributos in (
        ("w:tblW", {"w:w": str(sum(larguras)), "w:type": "dxa"}),
        ("w:tblLayout", {"w:type": "fixed"}),
    ):
        elemento = tblPr.find(qn(tag))
        if elemento is None:
            elemento = tblPr.makeelement(qn(tag), {})
            tblPr.append(elemento)
        for chave, valor in atributos.items():
            elemento.set(qn(chave), valor)

    # margem interna lateral menor (0,1 cm em vez dos 0,19 cm padrão do
    # Word) — libera ~4% de largura útil, o suficiente para CNPJ/NIT/
    # benefício caberem em uma linha sem estourar os 100%
    cell_mar = tblPr.find(qn("w:tblCellMar"))
    if cell_mar is None:
        cell_mar = tblPr.makeelement(qn("w:tblCellMar"), {})
        tblPr.append(cell_mar)
    for lado in ("w:left", "w:right"):
        elemento = cell_mar.find(qn(lado))
        if elemento is None:
            elemento = cell_mar.makeelement(qn(lado), {})
            cell_mar.append(elemento)
        elemento.set(qn("w:w"), "57")
        elemento.set(qn("w:type"), "dxa")

    for linha in tabela.rows:
        indice = 0
        for tc in linha._tr.findall(qn("w:tc")):
            tcPr = tc.find(qn("w:tcPr"))
            if tcPr is None:
                tcPr = tc.makeelement(qn("w:tcPr"), {})
                tc.insert(0, tcPr)
            grid_span = tcPr.find(qn("w:gridSpan"))
            span = int(grid_span.get(qn("w:val"))) if grid_span is not None else 1
            largura_celula = sum(larguras[indice:indice + span])
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = tcPr.makeelement(qn("w:tcW"), {})
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(largura_celula))
            tcW.set(qn("w:type"), "dxa")
            indice += span


def _mesclar_horizontais(tabela, indice_linha: int, spans) -> None:
    """Reproduz as mesclagens horizontais (gridSpan) de uma linha da tabela
    de origem. As células vinham preenchidas com o valor duplicado em cada
    coluna coberta; depois do merge, remove os parágrafos repetidos."""
    if not spans:
        return
    for col, span in enumerate(spans):
        if span and span > 1:
            celula = tabela.cell(indice_linha, col).merge(tabela.cell(indice_linha, col + span - 1))
            for paragrafo_extra in celula.paragraphs[1:]:
                paragrafo_extra._p.getparent().remove(paragrafo_extra._p)


def _mesclar_verticais_originais(tabela, indice_nome: int, linhas_dados: list[list[str]]) -> None:
    """Reproduz as mesclagens verticais que o documento original tinha nas
    DEMAIS colunas (ex.: "Processo", "Diferença entre DCB e DIB") — sem
    isso, o valor desembrulhado sai repetido em células separadas. A coluna
    do nome fica de fora: ela é tratada por `_mesclar_coluna_nome`, que
    agrupa por segurado. Só mescla quando a linha de baixo estava marcada
    como continuação (`vmerges`) E o valor confere — proteção para o caso
    de as linhas terem sido reagrupadas fora da ordem original."""
    if not linhas_dados:
        return
    n_colunas = len(linhas_dados[0])
    total = len(linhas_dados)
    for col in range(n_colunas):
        if col == indice_nome:
            continue
        i = 0
        while i < total:
            spans = getattr(linhas_dados[i], "spans", None)
            if spans and spans[col] == 0:
                i += 1
                continue
            j = i
            while j + 1 < total:
                proxima = linhas_dados[j + 1]
                vmerges = getattr(proxima, "vmerges", None)
                if not (vmerges and vmerges[col] and proxima[col] == linhas_dados[i][col]):
                    break
                j += 1
            if j > i:
                celula = tabela.cell(i + 1, col).merge(tabela.cell(j + 1, col))
                for paragrafo_extra in celula.paragraphs[1:]:
                    paragrafo_extra._p.getparent().remove(paragrafo_extra._p)
            i = j + 1


def _mesclar_coluna_nome(tabela, indice_nome: int, linhas_dados: list[list[str]]) -> None:
    """Mescla verticalmente as células da coluna de nome que pertencem ao
    mesmo segurado. `_Cell.merge()` do python-docx concatena os parágrafos
    das células mescladas (não deduplica), então cada linha já trazia o
    mesmo nome nas células desembrulhadas — depois do merge sobram
    parágrafos duplicados, que são removidos aqui mantendo só o primeiro
    (já com a fonte/estilo clonados aplicados antes do merge)."""
    total = len(linhas_dados)
    i = 0
    while i < total:
        nome = linhas_dados[i][indice_nome]
        j = i
        while j + 1 < total and linhas_dados[j + 1][indice_nome] == nome:
            j += 1
        if j > i:
            celula_mesclada = tabela.cell(i + 1, indice_nome).merge(tabela.cell(j + 1, indice_nome))
            for paragrafo_extra in celula_mesclada.paragraphs[1:]:
                paragrafo_extra._p.getparent().remove(paragrafo_extra._p)
        i = j + 1


def construir_tabela(doc: Document, cabecalhos: list[str], linhas_dados: list[list[str]],
                      indice_nome: int, tabela_referencia=None, larguras_colunas=None):
    por_copia = (
        tabela_referencia is not None
        and bool(linhas_dados)
        and all(getattr(linha, "tr", None) is not None for linha in linhas_dados)
    )
    if por_copia:
        tabela = _construir_tabela_por_copia(doc, tabela_referencia, linhas_dados)
    else:
        tabela = _construir_tabela_reconstruida(doc, cabecalhos, linhas_dados, indice_nome,
                                                tabela_referencia)

    if larguras_colunas:
        _aplicar_larguras_colunas(doc, tabela, larguras_colunas)

    # larguras oficiais do TABELAS.docx têm prioridade: as larguras do
    # arquivo importado variam de usuário para usuário
    larguras_oficiais = _larguras_do_catalogo(cabecalhos)
    if larguras_oficiais:
        _aplicar_larguras_twips(tabela, larguras_oficiais)

    _centralizar_tabela(tabela)

    if por_copia:
        _anexar_setas_flutuantes(doc, tabela, tabela_referencia, linhas_dados)
    return tabela


def _construir_tabela_reconstruida(doc: Document, cabecalhos: list[str], linhas_dados: list[list[str]],
                                    indice_nome: int, tabela_referencia=None):
    """Reconstrução célula a célula, usada quando as linhas não trazem o
    `w:tr` de origem (texto colado): clona só a aparência de uma linha de
    referência, então formatação pontual de runs não é preservada."""
    n_linhas = 1 + len(linhas_dados)
    n_colunas = len(cabecalhos)
    tabela = doc.add_table(rows=n_linhas, cols=n_colunas)

    linha_cabecalho_ref = None
    linha_corpo_ref = None
    if tabela_referencia is not None:
        try:
            # atribuir o objeto de estilo da tabela de origem (documento
            # diferente) pode não resolver corretamente; usar o nome faz o
            # python-docx procurar/criar o estilo dentro do próprio modelo.
            if tabela_referencia.style is not None:
                tabela.style = tabela_referencia.style.name
        except Exception:
            pass
        _clonar_largura_tabela(tabela_referencia, tabela)
        linha_cabecalho_ref = tabela_referencia.rows[0]
        if len(tabela_referencia.rows) > 1:
            linha_corpo_ref = tabela_referencia.rows[1]

    if linha_cabecalho_ref is not None:
        _clonar_altura_linha(linha_cabecalho_ref, tabela.rows[0])

    for col, texto in enumerate(cabecalhos):
        celula = tabela.rows[0].cells[col]
        celula.text = texto
        if linha_cabecalho_ref is not None and col < len(linha_cabecalho_ref.cells):
            ref_cell = linha_cabecalho_ref.cells[col]
            _clonar_aparencia_celula(ref_cell, celula)
            if ref_cell.paragraphs:
                _clonar_fonte_paragrafo(ref_cell.paragraphs[0], celula.paragraphs[0])
        # cabeçalho é sempre negrito e sempre a cor de destaque da tabela,
        # independente do que a célula de origem tinha clonado
        _forcar_negrito(celula)
        _sombrear_celula(celula, COR_TABELA_DESTAQUE)

    indice_item = _detectar_coluna_item(cabecalhos)
    for i, linha_dados in enumerate(linhas_dados, start=1):
        if linha_corpo_ref is not None:
            _clonar_altura_linha(linha_corpo_ref, tabela.rows[i])
        for col, valor in enumerate(linha_dados):
            celula = tabela.rows[i].cells[col]
            celula.text = valor
            if linha_corpo_ref is not None and col < len(linha_corpo_ref.cells):
                ref_cell = linha_corpo_ref.cells[col]
                _clonar_aparencia_celula(ref_cell, celula)
                if ref_cell.paragraphs:
                    _clonar_fonte_paragrafo(ref_cell.paragraphs[0], celula.paragraphs[0])
            if col == indice_item:
                _forcar_negrito(celula)
                _sombrear_celula(celula, COR_TABELA_DESTAQUE)

    _mesclar_horizontais(tabela, 0, getattr(cabecalhos, "spans", None))
    for i, linha_dados in enumerate(linhas_dados, start=1):
        _mesclar_horizontais(tabela, i, getattr(linha_dados, "spans", None))

    _mesclar_verticais_originais(tabela, indice_nome, linhas_dados)
    _mesclar_coluna_nome(tabela, indice_nome, linhas_dados)

    if tabela_referencia is None:
        _aplicar_estilo_padrao_grade(tabela)

    return tabela


def inserir_tabela_no_marcador(doc: Document, paragrafo_marcador: Paragraph, cabecalhos: list[str],
                                linhas_dados: list[list[str]], indice_nome: int, tabela_referencia=None,
                                larguras_colunas=None):
    tabela = construir_tabela(doc, cabecalhos, linhas_dados, indice_nome, tabela_referencia, larguras_colunas)
    # move a tabela e, junto, os parágrafos de formas flutuantes (setas)
    # que a construção por cópia possa ter anexado logo depois dela
    for elemento in [tabela._tbl] + list(getattr(tabela, "_paragrafos_setas", [])):
        elemento.getparent().remove(elemento)
        paragrafo_marcador._p.addprevious(elemento)

    # o parágrafo do marcador fica (não é removido, só esvaziado e
    # encolhido ao mínimo): no modelo padrão ele é o ÚLTIMO parágrafo do
    # corpo do documento, logo após a tabela. O Word sempre exige um
    # parágrafo depois de uma tabela nessa posição — quando a tabela já
    # preenche a página quase até o fim (tabelas com muitas linhas), até
    # um parágrafo vazio de tamanho normal não cabe mais no que sobrou e
    # transborda para uma página nova, inteiramente em branco. Encolher
    # esse parágrafo ao tamanho mínimo (fonte e entrelinha de 1pt, sem
    # espaçamento) evita esse transbordo.
    for run in list(paragrafo_marcador.runs):
        run.text = ""
    if not paragrafo_marcador.runs:
        paragrafo_marcador.add_run("")
    for run in paragrafo_marcador.runs:
        run.font.size = Pt(1)
    formato = paragrafo_marcador.paragraph_format
    formato.space_before = Pt(0)
    formato.space_after = Pt(0)
    formato.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    formato.line_spacing = Pt(1)
    return tabela


# ── Montagem das 3 variantes de página ────────────────────────────────────
#
# Um modelo pode ter dois formatos:
#   - simples: um único bloco {{TITULO}} + {{TABELA}}, reaproveitado como
#     está nas 3 variantes (capa sem tabela, tabela completa, por segurado).
#   - com capa própria: além disso, um bloco de capa separado, delimitado
#     pelo marcador {{TITULO_CAPA}} (título/logo/página de rosto, sem
#     tabela) que vem ANTES do bloco {{TITULO}}/{{TABELA}} no documento.
#     Cada variante usa só o bloco que faz sentido para ela e descarta o
#     outro, preservando o restante do layout (branding, fontes, cores).

def montar_capa(caminho_modelo: str, titulo: str, topico: str | None = None) -> Document:
    doc = Document(str(caminho_modelo))

    paragrafo_capa = localizar_paragrafo_marcador(doc, MARCADOR_TITULO_CAPA)
    if paragrafo_capa is not None:
        _substituir_no_paragrafo(paragrafo_capa, MARCADOR_TITULO_CAPA, titulo)
        if topico is not None:
            substituir_topico(doc, topico)
        paragrafo_titulo = localizar_paragrafo_marcador(doc, MARCADOR_TITULO)
        if paragrafo_titulo is not None:
            _remover_de(doc, paragrafo_titulo._p)
        return doc

    substituir_titulo(doc, titulo)
    marcador = localizar_paragrafo_marcador(doc)
    if marcador is not None:
        remover_paragrafo(marcador)
    return doc


def montar_tabela_completa(caminho_modelo: str, titulo: str, cabecalhos: list[str],
                            linhas: list[list[str]], indice_nome: int, tabela_referencia=None,
                            larguras_colunas=None) -> Document:
    doc = Document(str(caminho_modelo))

    paragrafo_capa = localizar_paragrafo_marcador(doc, MARCADOR_TITULO_CAPA)
    if paragrafo_capa is not None:
        _remover_ate(doc, paragrafo_capa._p)

    substituir_titulo(doc, titulo)
    marcador = localizar_paragrafo_marcador(doc)
    if marcador is None:
        raise ValueError('Marcador "{{TABELA}}" não encontrado no modelo.')
    inserir_tabela_no_marcador(doc, marcador, cabecalhos, linhas, indice_nome, tabela_referencia,
                               larguras_colunas)
    return doc


def montar_pagina_segurado(caminho_modelo: str, nome_segurado: str, linhas_segurado: list[list[str]],
                            cabecalhos: list[str], indice_nome: int, tabela_referencia=None,
                            larguras_colunas=None) -> Document:
    doc = Document(str(caminho_modelo))

    paragrafo_capa = localizar_paragrafo_marcador(doc, MARCADOR_TITULO_CAPA)
    if paragrafo_capa is not None:
        _remover_ate(doc, paragrafo_capa._p)

    substituir_titulo(doc, nome_segurado)
    marcador = localizar_paragrafo_marcador(doc)
    if marcador is None:
        raise ValueError('Marcador "{{TABELA}}" não encontrado no modelo.')
    inserir_tabela_no_marcador(doc, marcador, cabecalhos, linhas_segurado, indice_nome, tabela_referencia,
                               larguras_colunas)
    return doc


def _termina_com_quebra_pagina(doc: Document) -> bool:
    """A capa (bloco {{TITULO_CAPA}}) já termina com uma quebra de página
    manual herdada do modelo original (ela separava a capa da tabela no
    documento de origem) — sem essa checagem, `mesclar_documentos` somaria
    mais uma quebra em cima dessa e geraria uma página em branco."""
    ultimo = None
    for filho in doc.element.body:
        if filho.tag == qn("w:sectPr"):
            break
        ultimo = filho
    if ultimo is None or ultimo.tag != qn("w:p"):
        return False
    return any(br.get(qn("w:type")) == "page" for br in ultimo.iter(qn("w:br")))


def mesclar_documentos(documentos: list[Document]) -> Document:
    """Junta vários documentos (todos derivados do mesmo modelo.docx — mesmo
    cabeçalho/rodapé, estilos e tema) num único Document, um atrás do outro
    com quebra de página entre eles. Usa o primeiro como base e mantém só o
    `sectPr` dele (tamanho/margens/cabeçalho/rodapé da página) — como todos
    vêm do mesmo modelo, são idênticos entre si."""
    if not documentos:
        raise ValueError("Nenhum documento para mesclar.")

    base = documentos[0]
    sectPr_base = base.element.body.find(qn("w:sectPr"))

    for extra in documentos[1:]:
        if not _termina_com_quebra_pagina(base):
            base.add_page_break()
        for filho in list(extra.element.body):
            if filho.tag == qn("w:sectPr"):
                continue
            clone = copy.deepcopy(filho)
            if sectPr_base is not None:
                sectPr_base.addprevious(clone)
            else:
                base.element.body.append(clone)

    return base
