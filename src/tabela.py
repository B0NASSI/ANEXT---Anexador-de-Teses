# -*- coding: utf-8 -*-
"""
Leitura da tabela de segurados, de um .docx (com células mescladas
verticalmente entre benefícios do mesmo segurado) ou de texto colado no
formato de tabela do Word (colunas por tabulação, linhas por quebra de
linha). Em ambos os casos o resultado final é o mesmo: uma lista de linhas
"desembrulhadas" (nome preenchido em toda linha, sem mesclagem) agrupadas
por segurado, na ordem de primeira aparição.
"""

import unicodedata
from dataclasses import dataclass

from docx import Document
from docx.oxml.ns import qn

# radicais, casados por substring: cobrem todas as variações de gênero e
# número — EMPREGADA(S), EMPREGADO(A), EMPREGADOS (AS), SEGURADA(S) etc.
PRIORIDADE_COLUNA_NOME = ["empregad", "segurad", "nome"]

# tabelas sem coluna de segurado (ex.: tese de "Bloqueio de Rotatividade",
# que lista por empresa) são agrupadas por CNPJ (+ vigência, se houver)
PRIORIDADE_COLUNA_CNPJ = ["cnpj"]
PRIORIDADE_COLUNA_VIGENCIA = ["vigenc"]


@dataclass
class TabelaSegurados:
    cabecalhos: list[str]
    linhas: list[list[str]]
    indice_nome: int
    grupos: list[tuple[str, list[list[str]]]]


class LinhaComSpans(list):
    """Linha de células que carrega o gridSpan de cada posição da grade:
    1 = célula normal, N = célula mesclada sobre N colunas, 0 = posição
    coberta pela mesclagem anterior. Continua sendo uma list[str] comum
    (valores desembrulhados), então todo o código existente funciona igual;
    quem souber ler `spans` pode reproduzir as mesclagens horizontais.
    `vmerges` marca, por coluna, se a célula era continuação de uma
    mesclagem vertical no documento original (True) — permite reproduzir
    as mesclagens verticais de qualquer coluna, não só a do nome.
    `tr` guarda o elemento `w:tr` de origem da linha: com ele a montagem
    copia a linha inteira por deepcopy, preservando formatação de runs
    (negrito/cor pontuais) e formas flutuantes ancoradas nas células."""

    def __init__(self, valores, spans=None, vmerges=None, tr=None):
        super().__init__(valores)
        self.spans = spans
        self.vmerges = vmerges
        self.tr = tr


def _normalizar(texto: str) -> str:
    sem_acento = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode("ascii")
    return sem_acento.strip().lower()


def _detectar_coluna(cabecalhos: list[str], prioridades: list[str]) -> int | None:
    normalizados = [_normalizar(h) for h in cabecalhos]
    for alvo in prioridades:
        for i, cabecalho in enumerate(normalizados):
            if cabecalho == alvo:
                return i
    for alvo in prioridades:
        for i, cabecalho in enumerate(normalizados):
            if alvo in cabecalho:
                return i
    return None


def detectar_coluna_nome(cabecalhos: list[str]) -> int:
    indice = _detectar_coluna(cabecalhos, PRIORIDADE_COLUNA_NOME)
    if indice is None:
        raise ValueError(
            'Não foi possível identificar a coluna de nome do segurado na tabela '
            '(esperava um cabeçalho como "Empregados", "Segurado" ou "Nome").'
        )
    return indice


def _identificar_agrupamento(cabecalhos: list[str]) -> tuple[str, int, int | None]:
    """Decide como a tabela deve ser agrupada em páginas: pelo nome do
    segurado (padrão, ex.: tabelas de FAP) ou, quando não há essa coluna,
    por CNPJ (+ vigência, se houver) — caso da tese de "Bloqueio de
    Rotatividade", que lista por empresa em vez de por pessoa. Retorna
    (modo, índice da coluna principal, índice da coluna de vigência —
    só usado no modo "cnpj")."""
    indice_nome = _detectar_coluna(cabecalhos, PRIORIDADE_COLUNA_NOME)
    if indice_nome is not None:
        return "nome", indice_nome, None
    indice_cnpj = _detectar_coluna(cabecalhos, PRIORIDADE_COLUNA_CNPJ)
    if indice_cnpj is not None:
        return "cnpj", indice_cnpj, _detectar_coluna(cabecalhos, PRIORIDADE_COLUNA_VIGENCIA)
    raise ValueError(
        'Não foi possível identificar a coluna de agrupamento da tabela '
        '(esperava um cabeçalho como "Empregados"/"Segurado"/"Nome", ou "CNPJ").'
    )


def _agrupar_por_nome(linhas: list[list[str]], indice_nome: int) -> list[tuple[str, list[list[str]]]]:
    """Agrupa linhas consecutivas do MESMO segurado (vários benefícios um
    embaixo do outro) numa única página. Se o nome se repetir depois de
    outras linhas no meio — ex.: um benefício no item 2 e outro no item 7
    da mesma pessoa — trata como um NOVO grupo (página separada): esse
    espaçamento é sinal de que o usuário quis separar os benefícios, ao
    contrário de linhas seguidas, que continuam sendo juntadas como
    sempre."""
    grupos: list[tuple[str, list[list[str]]]] = []
    ultimo_nome = None
    for linha in linhas:
        nome = linha[indice_nome].strip()
        if not nome:
            continue
        if nome != ultimo_nome:
            grupos.append((nome, []))
            ultimo_nome = nome
        grupos[-1][1].append(linha)
    return grupos


def _agrupar_por_cnpj(
    linhas: list[list[str]], indice_cnpj: int, indice_vigencia: int | None,
) -> list[tuple[str, list[list[str]]]]:
    """Equivalente a `_agrupar_por_nome`, mas para tabelas sem segurado (ex.:
    Rotatividade): agrupa linhas consecutivas do mesmo CNPJ+vigência, com o
    rótulo da página no formato "VIGÊNCIA {ano} – CNPJ Nº {cnpj}"."""
    grupos: list[tuple[str, list[list[str]]]] = []
    ultima_chave = None
    for linha in linhas:
        cnpj = linha[indice_cnpj].strip()
        vigencia = linha[indice_vigencia].strip() if indice_vigencia is not None else ""
        if not cnpj:
            continue
        chave = (cnpj, vigencia)
        if chave != ultima_chave:
            rotulo = f"VIGÊNCIA {vigencia} – CNPJ Nº {cnpj}" if vigencia else f"CNPJ Nº {cnpj}"
            grupos.append((rotulo, []))
            ultima_chave = chave
        grupos[-1][1].append(linha)
    return grupos


def _agrupar(
    modo: str, linhas: list[list[str]], indice_principal: int, indice_vigencia: int | None,
) -> list[tuple[str, list[list[str]]]]:
    if modo == "nome":
        return _agrupar_por_nome(linhas, indice_principal)
    return _agrupar_por_cnpj(linhas, indice_principal, indice_vigencia)


# ── Importação de .docx ──────────────────────────────────────────────────

def _vmerge_valor(celula) -> str | None:
    tcPr = celula._tc.tcPr
    if tcPr is None:
        return None
    vmerge = tcPr.find(qn("w:vMerge"))
    if vmerge is None:
        return None
    return vmerge.get(qn("w:val")) or "continue"


def _texto_celula(celula) -> str:
    return "\n".join(p.text for p in celula.paragraphs).strip()


def _spans_da_linha(celulas, n_colunas: int) -> list[int]:
    """Detecta mesclagens horizontais: `row.cells` repete o MESMO objeto de
    célula para cada coluna da grade coberta por um gridSpan — a identidade
    do `_tc` revela onde uma célula começa e quantas colunas ela cobre."""
    spans: list[int] = []
    i = 0
    limite = min(len(celulas), n_colunas)
    while i < limite:
        tc = celulas[i]._tc
        span = 1
        while i + span < limite and celulas[i + span]._tc is tc:
            span += 1
        spans.append(span)
        spans.extend([0] * (span - 1))
        i += span
    while len(spans) < n_colunas:
        spans.append(1)
    return spans


def extrair_de_docx(caminho: str) -> tuple[TabelaSegurados, "docx.table.Table"]:
    doc = Document(str(caminho))
    if not doc.tables:
        raise ValueError("O documento selecionado não contém nenhuma tabela.")
    tabela_word = max(doc.tables, key=lambda t: len(t.rows))

    linhas_grid = tabela_word.rows
    if len(linhas_grid) < 2:
        raise ValueError("A tabela encontrada não tem linhas de dados (só cabeçalho).")

    celulas_cabecalho = linhas_grid[0].cells
    cabecalhos = LinhaComSpans(
        [_texto_celula(c) for c in celulas_cabecalho],
        _spans_da_linha(celulas_cabecalho, len(celulas_cabecalho)),
    )
    n_colunas = len(cabecalhos)

    ultimo_valor_coluna: list[str | None] = [None] * n_colunas
    # `row.cells` devolve a célula do TOPO da mesclagem para as linhas de
    # continuação vertical (o texto já vem "desembrulhado") — a continuação
    # é detectada comparando a identidade do `_tc` com a linha anterior,
    # mesma técnica dos gridSpans horizontais.
    tc_linha_anterior: list = [None] * n_colunas
    linhas: list[list[str]] = []
    for linha_word in linhas_grid[1:]:
        celulas = linha_word.cells
        linha_valores = []
        continuacoes = []
        tcs_atuais: list = [None] * n_colunas
        for i in range(n_colunas):
            if i >= len(celulas):
                linha_valores.append(ultimo_valor_coluna[i] or "")
                continuacoes.append(False)
                continue
            celula = celulas[i]
            tcs_atuais[i] = celula._tc
            if celula._tc is tc_linha_anterior[i] or _vmerge_valor(celula) == "continue":
                valor = ultimo_valor_coluna[i] or _texto_celula(celula)
                continuacoes.append(True)
            else:
                valor = _texto_celula(celula)
                ultimo_valor_coluna[i] = valor
                continuacoes.append(False)
            linha_valores.append(valor)
        tc_linha_anterior = tcs_atuais
        if any(v.strip() for v in linha_valores):
            linhas.append(LinhaComSpans(linha_valores, _spans_da_linha(celulas, n_colunas), continuacoes,
                                        linha_word._tr))

    modo, indice_principal, indice_vigencia = _identificar_agrupamento(cabecalhos)
    grupos = _agrupar(modo, linhas, indice_principal, indice_vigencia)
    tabela_segurados = TabelaSegurados(
        cabecalhos=cabecalhos, linhas=linhas, indice_nome=indice_principal, grupos=grupos,
    )
    return tabela_segurados, tabela_word


# ── Texto colado (formato de tabela do Word) ─────────────────────────────

def extrair_de_texto(texto: str) -> TabelaSegurados:
    linhas_brutas = [
        linha for linha in texto.replace("\r\n", "\n").replace("\r", "\n").split("\n")
        if linha.strip() != ""
    ]
    if len(linhas_brutas) < 2:
        raise ValueError(
            "Cole o cabeçalho e ao menos uma linha de dados (colunas separadas por tabulação)."
        )

    tabela_bruta = [linha.split("\t") for linha in linhas_brutas]
    n_colunas = max(len(linha) for linha in tabela_bruta)
    tabela_bruta = [linha + [""] * (n_colunas - len(linha)) for linha in tabela_bruta]

    cabecalhos = [celula.strip() for celula in tabela_bruta[0]]
    modo, indice_principal, indice_vigencia = _identificar_agrupamento(cabecalhos)

    linhas: list[list[str]] = []
    ultimo_valor_principal = None
    for linha_bruta in tabela_bruta[1:]:
        linha = [celula.strip() for celula in linha_bruta]
        if not any(linha):
            continue
        if not linha[indice_principal] and ultimo_valor_principal:
            linha[indice_principal] = ultimo_valor_principal
        else:
            ultimo_valor_principal = linha[indice_principal]
        linhas.append(linha)

    grupos = _agrupar(modo, linhas, indice_principal, indice_vigencia)
    return TabelaSegurados(
        cabecalhos=cabecalhos, linhas=linhas, indice_nome=indice_principal, grupos=grupos,
    )
